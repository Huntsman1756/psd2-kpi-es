"""CaixaBank PSD2 statistics parser.

Sources (verified 2026-10):
  - current quarters: XLSX 'CaixaBank-<yyyy>Q<q>-InformeTrimestralPSD2.xlsx'
    linked from the API Store page. Contains a 'Disponibilidad & rendimiento'
    sheet with daily rows (raw numeric values) and a 'Leyenda' sheet with KPI
    definitions.
  - older quarters: one-page PDFs 'caixabank-<yyyy>q<q>-psd2.pdf' hosted on
    imagin.com, discovered by a bounded URL-pattern probe.

Column semantics (documented in the file's 'Leyenda' sheet):
  API/Home Banking each: Disponibilidad | AIS{TMR ms, %Éxito} |
                         PIS{TMR ms, %Éxito, %Éxito consolidación}
  '% Éxito consolidación' is a *different* metric (transfer consolidation
  success) and gets its own comparability_group.
  'Faltan datos' cells => not reported (skipped).

Comparability note (from 'Leyenda'): API availability follows the Delegated
Regulation criterion while Home Banking availability follows internal
criteria. Same comparability_group is therefore PARTIAL at best.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime

import openpyxl
import pdfplumber

from psd2_kpi_es.models import (
    Aggregation,
    Comparability,
    Entity,
    InterfaceType,
    Metric,
    Observation,
    ParseResult,
    ParseWarning,
    PeriodType,
    Service,
    SourceArtifact,
    Unit,
)
from psd2_kpi_es.parsers.base import quarter_bounds
from psd2_kpi_es.parsers.textnorm import parse_es_date, parse_es_number, parse_es_percent

NAME = "caixabank"
VERSION = "caixabank:1"

_INDEX_XLSX_RE = re.compile(r"[^\"'\s]*InformeTrimestralPSD2[^\"'\s]*\.xlsx", re.IGNORECASE)
_XLSX_Q_RE = re.compile(r"(\d{4})Q([1-4])", re.IGNORECASE)
_PDF_Q_RE = re.compile(r"caixabank-(\d{4})q([1-4])-psd2\.pdf", re.IGNORECASE)
_IMAGIN_PATTERN = "https://www.imagin.com/documents/554420/699403/caixabank-{y}q{q}-psd2.pdf"
_DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")

# earliest quarter ever observed on imagin.com (probed 2026-10: 2024q2 exists)
_IMAGIN_FIRST = (2024, 2)

_DEF_API_DISP = (
    "Tiempo en línea diario, calculada según el criterio del Reglamento "
    "Técnico Delegado (RTS 2018/389) per publisher's Leyenda."
)
_DEF_HB_DISP = (
    "Tiempo en línea diario calculado siguiendo criterios de disponibilidad "
    "internos del publisher's Leyenda."
)
_DEF_TMR = "Tiempo medio de respuesta en ms sobre las operaciones OK."
_DEF_EXITO = "Porcentaje de peticiones procesadas correctamente sobre el volumen total."
_DEF_CONS = (
    "Porcentaje de consolidaciones de transferencias completadas sobre el "
    "total de consolidaciones ejecutadas."
)


def discover_links(index_html: bytes) -> list[str]:
    text = index_html.decode("utf-8", errors="replace")
    urls = set()
    for m in _INDEX_XLSX_RE.findall(text):
        if m.startswith("http"):
            urls.add(m)
        elif m.startswith("/"):
            urls.add("https://www.caixabank.es" + m)

    # bounded pattern probe for historical imagin.com PDFs
    today = date.today()
    cy, cq = today.year, (today.month - 1) // 3
    y, q = _IMAGIN_FIRST
    while (y, q) <= (cy, cq):
        urls.add(_IMAGIN_PATTERN.format(y=y, q=q))
        q += 1
        if q == 5:
            y, q = y + 1, 1
    return sorted(urls)


def period_hint(url: str) -> str | None:
    m = _XLSX_Q_RE.search(url) or _PDF_Q_RE.search(url)
    return f"{m.group(1)}Q{int(m.group(2))}" if m else None


def _obs(
    *,
    entity,
    artifact,
    day_start,
    day_end,
    ptype,
    plabel,
    interface,
    service,
    metric,
    value,
    unit,
    aggregation,
    group,
    raw_label,
    raw_value,
    raw_unit,
    locator,
    definition=None,
) -> Observation:
    return Observation(
        entity_id=entity.entity_id,
        entity_name=entity.legal_name,
        period_start=day_start,
        period_end=day_end,
        period_type=ptype,
        period_label=plabel,
        interface_type=interface,
        service=service,
        metric=metric,
        value=value,
        unit=unit,
        aggregation=aggregation,
        metric_definition=definition,
        comparability=Comparability.PARTIAL,
        comparability_group=group,
        source_id=artifact.source_id,
        source_url=artifact.source_url,
        source_document=locator,
        retrieved_at=artifact.retrieved_at,
        source_sha256=artifact.sha256,
        parser_name=NAME,
        parser_version=VERSION,
        raw_label=raw_label,
        raw_value=raw_value,
        raw_unit=raw_unit,
    )


def _emit_row(
    out, *, entity, artifact, start, end, ptype, plabel, vals, locator, raw_fmt=str, pct_scale=1.0
):
    """vals: 12 raw values [disp_api, ais_tmr, ais_ex, pis_tmr, pis_ex, pis_cons,
    disp_hb, ais_tmr, ais_ex, pis_tmr, pis_ex, pis_cons] or None for missing.
    pct_scale converts ratio values (xlsx) to percent (pdf supplies percent)."""
    if len(vals) != 12:
        raise ValueError(f"expected 12 cells, got {len(vals)}")

    def emit(idx, iface, service, metric, unit, group, raw_label, raw_unit, definition, scale=1.0):
        raw = vals[idx]
        if raw is None:
            return
        v = raw * scale if isinstance(raw, (int, float)) else raw
        out.append(
            _obs(
                entity=entity,
                artifact=artifact,
                day_start=start,
                day_end=end,
                ptype=ptype,
                plabel=plabel,
                interface=iface,
                service=service,
                metric=metric,
                value=v,
                unit=unit,
                aggregation=(
                    Aggregation.DAILY_MEAN if ptype is PeriodType.DAY else Aggregation.MEAN
                ),
                group=group,
                raw_label=raw_label,
                raw_value=raw_fmt(raw),
                raw_unit=raw_unit,
                locator=locator,
                definition=definition,
            )
        )

    for off, iface in ((0, InterfaceType.DEDICATED_API), (6, InterfaceType.PSU_INTERFACE)):
        disp_def = _DEF_API_DISP if iface is InterfaceType.DEDICATED_API else _DEF_HB_DISP
        emit(
            off + 0,
            iface,
            Service.ALL,
            Metric.AVAILABILITY,
            Unit.PERCENT,
            "availability_daily_pct",
            "Disponibilidad",
            "%",
            disp_def,
            scale=pct_scale,
        )
        emit(
            off + 1,
            iface,
            Service.AIS,
            Metric.RESPONSE_TIME,
            Unit.MS,
            "response_time_daily_mean_ms",
            "TMR (ms)",
            "ms",
            _DEF_TMR,
        )
        emit(
            off + 2,
            iface,
            Service.AIS,
            Metric.REQUEST_SUCCESS_RATE,
            Unit.PERCENT,
            "success_rate_daily_pct",
            "% Éxito",
            "%",
            _DEF_EXITO,
            scale=pct_scale,
        )
        emit(
            off + 3,
            iface,
            Service.PIS,
            Metric.RESPONSE_TIME,
            Unit.MS,
            "response_time_daily_mean_ms",
            "TMR (ms)",
            "ms",
            _DEF_TMR,
        )
        emit(
            off + 4,
            iface,
            Service.PIS,
            Metric.REQUEST_SUCCESS_RATE,
            Unit.PERCENT,
            "success_rate_daily_pct",
            "% Éxito",
            "%",
            _DEF_EXITO,
            scale=pct_scale,
        )
        emit(
            off + 5,
            iface,
            Service.PIS,
            Metric.REQUEST_SUCCESS_RATE,
            Unit.PERCENT,
            "success_rate_consolidation_pct",
            "% Éxito consolidación",
            "%",
            _DEF_CONS,
            scale=pct_scale,
        )


def _parse_xlsx(artifact, content, entity) -> ParseResult:
    warnings: list[ParseWarning] = []
    out: list[Observation] = []
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = None
    for name in wb.sheetnames:
        if "rendimiento" in name.lower() or "disponibilidad" in name.lower():
            ws = wb[name]
    if ws is None:
        return ParseResult(
            warnings=[ParseWarning(code="NO_SHEET", message=f"sheets: {wb.sheetnames}")]
        )

    def cv(cell):
        v = cell.value
        if v is None or (isinstance(v, str) and "faltan" in v.lower()):
            return None
        if isinstance(v, str):
            return float(parse_es_number(v))
        return float(v)

    for row in ws.iter_rows(min_row=1):
        first = row[0].value
        if isinstance(first, datetime):
            day = first.date()
            vals = [cv(c) for c in row[1:13]]
            _emit_row(
                out,
                entity=entity,
                artifact=artifact,
                start=day,
                end=day,
                ptype=PeriodType.DAY,
                plabel=day.isoformat(),
                vals=vals,
                locator=f"sheet '{ws.title}'",
                raw_fmt=lambda v: f"{v!r}",
                pct_scale=100.0,
            )
        elif isinstance(first, str) and first.strip().upper() == "TOTAL":
            vals = [cv(c) for c in row[1:13]]
            # period bounds from the filename hint
            hint = artifact.period_hint or ""
            m = re.fullmatch(r"(\d{4})Q([1-4])", hint)
            if m:
                qs, qe = quarter_bounds(int(m.group(1)), int(m.group(2)))
            else:
                days = [o.period_start for o in out]
                qs, qe = min(days), max(days)
                hint = f"{qs.year}Q{(qs.month - 1) // 3 + 1}"
            _emit_row(
                out,
                entity=entity,
                artifact=artifact,
                start=qs,
                end=qe,
                ptype=PeriodType.QUARTER,
                plabel=hint,
                vals=vals,
                locator=f"sheet '{ws.title}' TOTAL",
                raw_fmt=lambda v: f"{v!r}",
                pct_scale=100.0,
            )
    return ParseResult(observations=out, warnings=warnings)


_TOK_RE = re.compile(r"faltan datos|sin datos|[\d.,]+%?", re.IGNORECASE)


def _pdf_tokens(rest: str) -> list[float | None]:
    vals: list[float | None] = []
    for t in _TOK_RE.findall(rest):
        if t.lower() in ("faltan datos", "sin datos"):
            vals.append(None)
        else:
            vals.append(
                float(parse_es_percent(t)) if t.endswith("%") else float(parse_es_number(t))
            )
    return vals


def _parse_pdf(artifact, content, entity) -> ParseResult:
    """Two PDF layouts seen in the wild:

    a) single page, one row = date + 12 cells (API 6 + Home Banking 6);
    b) split table: block A rows = date + 7 cells (API 6 + HB Disponibilidad),
       block B rows = 5 cells (remaining HB metrics, no date), zipped by order.
    """
    warnings: list[ParseWarning] = []
    out: list[Observation] = []
    pending_a: list[tuple[tuple, list]] = []  # ((ptype,plabel,start,end,page), vals7)
    block_b: list[list] = []

    hint = artifact.period_hint or ""
    hm = re.fullmatch(r"(\d{4})Q([1-4])", hint)

    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page_no, page in enumerate(pdf.pages):
            for line in (page.extract_text() or "").splitlines():
                s = line.strip()
                m = re.match(r"^(\d{2}/\d{2}/\d{4})\s+(.*)$", s)
                is_total = s.upper().startswith("TOTAL")
                if m:
                    day = date.fromisoformat(parse_es_date(m.group(1)))
                    key = (PeriodType.DAY, day.isoformat(), day, day, page_no)
                    vals = _pdf_tokens(m.group(2))
                elif is_total:
                    if not hm:
                        warnings.append(
                            ParseWarning(
                                code="NO_TOTAL_PERIOD", message="TOTAL row without period hint"
                            )
                        )
                        continue
                    start, end = quarter_bounds(int(hm.group(1)), int(hm.group(2)))
                    key = (PeriodType.QUARTER, hint, start, end, page_no)
                    vals = _pdf_tokens(s[len("TOTAL") :])
                else:
                    vals = _pdf_tokens(s)
                    if len(vals) == 5:
                        block_b.append(vals)
                    continue

                if len(vals) == 12:
                    _emit_row(
                        out,
                        entity=entity,
                        artifact=artifact,
                        start=key[2],
                        end=key[3],
                        ptype=key[0],
                        plabel=key[1],
                        vals=vals,
                        locator=f"page {page_no}",
                        raw_fmt=str,
                        pct_scale=1.0,
                    )
                elif len(vals) == 7:
                    pending_a.append((key, vals))
                else:
                    warnings.append(
                        ParseWarning(
                            code="ROW_SHAPE",
                            message=f"{key[1]}: {len(vals)} cells (expected 7 or 12)",
                        )
                    )

    if pending_a:
        if len(block_b) != len(pending_a):
            warnings.append(
                ParseWarning(
                    code="BLOCK_MISMATCH",
                    message=f"{len(pending_a)} dated rows vs {len(block_b)} "
                    f"undated HB rows; HB metrics beyond Disponibilidad "
                    f"skipped for mismatched tail",
                )
            )
        for i, (key, vals) in enumerate(pending_a):
            ptype, plabel, start, end, page_no = key
            merged = list(vals)
            if i < len(block_b):
                merged += block_b[i]
            else:
                merged += [None] * 5
            _emit_row(
                out,
                entity=entity,
                artifact=artifact,
                start=start,
                end=end,
                ptype=ptype,
                plabel=plabel,
                vals=merged,
                locator=f"page {page_no}",
                raw_fmt=str,
                pct_scale=1.0,
            )

    return ParseResult(observations=out, warnings=warnings)


def parse(artifact: SourceArtifact, content: bytes, entity: Entity) -> ParseResult:
    fn = (artifact.filename or artifact.source_url).lower()
    if fn.endswith(".xlsx") or (artifact.content_type and "spreadsheet" in artifact.content_type):
        return _parse_xlsx(artifact, content, entity)
    if fn.endswith(".pdf") or content[:5] == b"%PDF-":
        return _parse_pdf(artifact, content, entity)
    return ParseResult(warnings=[ParseWarning(code="UNSUPPORTED", message=f"unknown format: {fn}")])


name = NAME
version = VERSION
